"""schedule_probe on real Google pages. No network."""
import os
import sys
import unittest
from datetime import date

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import schedule_probe as sp  # noqa: E402
from test_pipeline import fixture  # noqa: E402

BLR = sp.Place("bangalore", "BLR", "Bengaluru", "India")
KUL = sp.Place("kuala-lumpur", "KUL", "Kuala Lumpur", "Malaysia")
WEEK = [date(2026, 11, 2 + i) for i in range(7)]


class NonstopFlights(unittest.TestCase):
    def test_reads_nonstops_from_a_real_page(self):
        got = sp.nonstop_flights(fixture("blr_kul_nonstop"), "BLR", "KUL")
        self.assertIn("AirAsia", {f["airline"] for f in got})
        self.assertTrue(all(f["depart"] for f in got))
        # Same airline at the same time is one flight, however Google repeats it.
        self.assertEqual(len(got), len({(f["airline"], f["depart"]) for f in got}))

    def test_nonstops_in_the_best_flights_block_count(self):
        # The capped page lists the non-stops under "best flights" only.
        got = sp.nonstop_flights(fixture("blr_kul_2stop"), "BLR", "KUL")
        self.assertEqual({f["airline"] for f in got}, {"AirAsia", "IndiGo", "Malaysia Airlines"})
        self.assertIn("00:20", [f["depart"] for f in got])

    def test_no_service_is_an_answer(self):
        self.assertEqual(sp.nonstop_flights(fixture("pnq_hkt_nonstop_empty"), "PNQ", "HKT"), [])

    def test_captcha_is_not_an_answer(self):
        with self.assertRaises(ValueError):
            sp.nonstop_flights(fixture("captcha"), "BLR", "KUL")


class Summary(unittest.TestCase):
    def week(self, pattern):
        aa = [{"airline": "AirAsia", "depart": "23:25", "arrive": "06:20", "durationMinutes": 265, "aircraft": None}]
        return {d.isoformat(): (aa if c == "y" else None if c == "?" else []) for d, c in zip(WEEK, pattern)}

    def test_frequencies(self):
        cases = {"yyyyyyy": "Daily", "yyyyyy-": "5-6x weekly", "y-y-y--": "Alternate days (3-4x weekly)",
                 "y------": "1-2x weekly", "-------": "No direct"}
        for pattern, want in cases.items():
            rec = sp.summarise(BLR, KUL, WEEK, self.week(pattern))
            self.assertEqual(rec["frequency"], want, pattern)
            self.assertEqual(rec["direct"], "y" in pattern)

    def test_unread_days_are_not_counted_as_no_flight(self):
        rec = sp.summarise(BLR, KUL, WEEK, self.week("yyy?yyy"))
        self.assertEqual(rec["daysChecked"], 6)
        self.assertTrue(rec["frequency"].startswith("Incomplete"))
        self.assertIsNone(rec["days"]["Thu"])

    def test_airline_days(self):
        rec = sp.summarise(BLR, KUL, WEEK, self.week("y-y-y--"))
        self.assertEqual(rec["airlines"], {"AirAsia": ["Mon", "Wed", "Fri"]})
        self.assertEqual(rec["departures"], {"AirAsia": ["23:25"]})

    def test_week_starts_on_a_monday_three_weeks_out(self):
        start = sp.default_week_start(date(2026, 10, 5))
        self.assertEqual(start.weekday(), 0)
        self.assertGreaterEqual((start - date(2026, 10, 5)).days, 21)


class SheetFormulas(unittest.TestCase):
    """The sheet's formulas must agree with the script's own summary."""

    def test_formulas_match_python(self):
        try:
            from openpyxl import load_workbook
            from pycel import ExcelCompiler
        except ImportError:
            self.skipTest("needs openpyxl and pycel")
        import tempfile
        aa = [{"airline": "AirAsia", "depart": "23:25", "arrive": None, "durationMinutes": None, "aircraft": None}]
        pats = {"KUL": "yyyyyyy", "HKT": "y-y-y--", "FCO": "-------", "DPS": "yy?yyyy",
                "MLE": "yyyyyy-", "SIN": "y------", "CMB": "---?---"}
        recs = [sp.summarise(BLR, sp.Place(k.lower(), k, k, "X"), WEEK,
                             {d.isoformat(): (aa if c == "y" else None if c == "?" else []) for d, c in zip(WEEK, p)})
                for k, p in pats.items()]
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "s.xlsx")
            sp.write_xlsx(recs, path)
            ws = load_workbook(path)["Routes"]
            xl = ExcelCompiler(filename=path)
            for r in range(2, ws.max_row + 1):
                rec = next(x for x in recs if x["dest"] == ws.cell(r, 4).value)
                want = ("Yes" if rec["direct"] else "Unknown" if rec["daysChecked"] < 7 else "No",
                        rec["daysWithNonstop"],
                        "Incomplete" if rec["daysChecked"] < 7 else rec["frequency"])
                got = tuple(xl.evaluate(f"Routes!{c}{r}") for c in "FOP")
                self.assertEqual(got, want, rec["dest"])


if __name__ == "__main__":
    unittest.main()
