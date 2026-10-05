#!/usr/bin/env python3
"""Which routes have non-stop flights, on which weekdays, and by whom.

For every origin x destination in routes.json, asks Google Flights for
non-stop one-way flights on each day of one week, and writes:

  out/schedules.xlsx   the route sheet: direct yes/no, airlines, Mon..Sun,
                       daily / alternate days / weekly, departure times
  out/schedules.json   the same, for the scraper to sample only days that
                       actually have flights

One week of non-stop searches answers the schedule question because airline
schedules repeat weekly. The week sampled is a few weeks out by default, far
enough that flights are not sold out and near enough that the schedule is
already loaded.

Cost: 7 requests per route, ~190 routes, ~1,330 requests — about 2.5 hours at
the scraper's pace. Progress is saved after every route; re-running resumes.

Run it somewhere other than the droplet if you can (a laptop is ideal): it is
the same kind of traffic as the fare pass, and two streams from one IP is how
CAPTCHAs start. On the droplet it refuses to run while a fare pass holds the
lock.

Usage:
    python3 schedule_probe.py                         # every route
    python3 schedule_probe.py --origins bangalore     # one origin
    python3 schedule_probe.py --week-start 2026-11-02 # a specific Monday
    python3 schedule_probe.py --xlsx-only             # rebuild the sheet from the JSON
"""
from __future__ import annotations

import argparse
import json
import os
import random
import sys
import time
from datetime import date, timedelta
from typing import Any

import scrape
from scrape import DELAY, Place

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "out")
JSON_PATH = os.path.join(OUT, "schedules.json")
XLSX_PATH = os.path.join(OUT, "schedules.xlsx")
LOCK = "/var/lock/fk-flight-finder.lock"

DAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]

# CAPTCHA handling: wait, then try the same day again; after this many waits
# in a row, stop and keep what has been saved.
CAPTCHA_WAIT = 600
CAPTCHA_GIVE_UP = 3


def default_week_start(today: date | None = None) -> date:
    """The Monday at least three weeks out."""
    today = today or date.today()
    start = today + timedelta(days=21)
    return start + timedelta(days=(7 - start.weekday()) % 7)


def nonstop_flights(html: str, origin: str, dest: str) -> list[dict[str, Any]]:
    """Non-stop flights origin->dest on a one-way results page.

    Prices are ignored: a flight Google lists without a price still flies, and
    the schedule is the point here. Both result slots are read — the library
    reads only the second ("other flights"); the first holds "best flights"
    when Google splits the list.

    Raises ValueError when the page has no data block (a block, not an answer).
    Returns [] when Google answered with no flights.
    """
    from selectolax.lexbor import LexborHTMLParser

    node = LexborHTMLParser(html).css_first(r"script.ds\:1")
    if node is None:
        raise ValueError("no data block in response")
    raw = node.text().split("data:", 1)[1].rsplit(",", 1)[0]
    if raw.endswith("errorHasStatus: true"):
        return []
    payload = json.loads(raw)

    entries: list[Any] = []
    for slot in (2, 3):
        block = payload[slot] if len(payload) > slot else None
        if block and block[0]:
            entries.extend(block[0])

    if not entries and len(payload) < scrape.COMPLETE_PAYLOAD_MIN:
        # Sent before Google had results: "unknown", not "no flights".
        raise scrape.Unanswered(f"no results in a {len(payload)}-slot payload")

    seen: set[tuple[str, str | None]] = set()
    out: list[dict[str, Any]] = []
    for entry in entries:
        try:
            flight = entry[0]
            segs = flight[2]
            if len(segs) != 1:
                continue
            sf = segs[0]
            if sf[3] != origin or sf[6] != dest:
                continue
            airline = (flight[1] or ["Unknown"])[0]
            dep = scrape._clock(sf[8])
            key = (airline, dep)
            if key in seen:
                continue
            seen.add(key)
            out.append({
                "airline": airline,
                "depart": dep,
                "arrive": scrape._clock(sf[10]),
                "durationMinutes": sf[11] if isinstance(sf[11], int) else None,
                "aircraft": sf[17] if len(sf) > 17 and isinstance(sf[17], str) else None,
            })
        except (IndexError, TypeError):
            continue
    return sorted(out, key=lambda f: (f["depart"] or "", f["airline"]))


def _query(origin: str, dest: str, day: str) -> Any:
    ff = scrape.ff
    return ff.create_query(
        flights=[ff.FlightQuery(date=day, from_airport=origin, to_airport=dest)],
        trip="one-way",
        seat="economy",
        passengers=ff.Passengers(adults=1),
        currency="INR",
        language="en-US",
        max_stops=0,
    )


def probe_day(origin: str, dest: str, day: str) -> list[dict[str, Any]] | None:
    """Non-stop flights that day, or None when Google would not say."""
    captchas = 0
    for attempt in range(1 + CAPTCHA_GIVE_UP):
        try:
            html = scrape._fetch_html(_query(origin, dest, day), None)
        except Exception as exc:  # noqa: BLE001 — timeout, reset: one retry
            print(f"      ! {origin}->{dest} {day}: {type(exc).__name__}", flush=True)
            if attempt:
                return None
            time.sleep(DELAY)
            continue
        if scrape._is_captcha(html):
            captchas += 1
            if captchas > CAPTCHA_GIVE_UP:
                raise scrape.Blocked()
            print(f"      ! CAPTCHA; waiting {CAPTCHA_WAIT // 60} min", flush=True)
            time.sleep(CAPTCHA_WAIT)
            continue
        try:
            return nonstop_flights(html, origin, dest)
        except (ValueError, IndexError, TypeError, json.JSONDecodeError) as exc:
            print(f"      ! {origin}->{dest} {day}: unreadable ({type(exc).__name__})", flush=True)
            if attempt:
                return None
            time.sleep(DELAY)
    return None


def summarise(origin: Place, dest: Place, week: list[date], by_day: dict[str, Any]) -> dict[str, Any]:
    """One route's week as a record: which days, which airlines, how often."""
    days: dict[str, bool | None] = {}
    airline_days: dict[str, list[str]] = {}
    departures: dict[str, set[str]] = {}
    for d in week:
        name = DAYS[d.weekday()]
        flights = by_day.get(d.isoformat())
        if flights is None:
            days[name] = None
            continue
        days[name] = bool(flights)
        for f in flights:
            airline_days.setdefault(f["airline"], [])
            if name not in airline_days[f["airline"]]:
                airline_days[f["airline"]].append(name)
            if f["depart"]:
                departures.setdefault(f["airline"], set()).add(f["depart"])
    known = [v for v in days.values() if v is not None]
    flying = sum(1 for v in known if v)
    return {
        "origin": origin.code,
        "originSlug": origin.slug,
        "originCity": origin.city,
        "dest": dest.code,
        "destSlug": dest.slug,
        "destCity": dest.city,
        "destCountry": dest.country,
        "weekStart": week[0].isoformat(),
        "days": days,
        "daysChecked": len(known),
        "daysWithNonstop": flying,
        "direct": flying > 0,
        "frequency": frequency(flying, len(known)),
        "airlines": {a: [n for n in DAYS if n in ds] for a, ds in sorted(airline_days.items())},
        "departures": {a: sorted(t) for a, t in sorted(departures.items())},
        "flightsByDay": by_day,
    }


def frequency(flying: int, checked: int) -> str:
    if checked < 7:
        return f"Incomplete ({checked}/7 days read)" if flying else f"Unknown ({checked}/7 days read)"
    if flying == 7:
        return "Daily"
    if flying >= 5:
        return "5-6x weekly"
    if flying >= 3:
        return "Alternate days (3-4x weekly)"
    if flying >= 1:
        return "1-2x weekly"
    return "No direct"


def _days_label(names: list[str]) -> str:
    return "Daily" if len(names) == 7 else ", ".join(names)


def write_xlsx(records: list[dict[str, Any]], path: str) -> None:
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    font = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", start_color="1F4E78")
    green = PatternFill("solid", start_color="C6EFCE")
    grey = PatternFill("solid", start_color="EDEDED")

    wb = Workbook()
    ws = wb.active
    ws.title = "Routes"
    cols = ["Origin", "From", "Destination", "To", "Country", "Direct", "Airlines (days)",
            *DAYS, "Days / week", "Frequency", "Non-stop departures", "Week checked", "Note"]
    ws.append(cols)
    first_day = cols.index("Mon") + 1
    last_day = first_day + 6
    c_days = cols.index("Days / week") + 1
    c_direct = cols.index("Direct") + 1
    c_freq = cols.index("Frequency") + 1

    for rec in sorted(records, key=lambda r: (r["originCity"], not r["direct"], r["destCity"])):
        r = ws.max_row + 1
        airlines = "; ".join(f"{a} ({_days_label(ds)})" for a, ds in rec["airlines"].items())
        deps = "; ".join(f"{a} {', '.join(t)}" for a, t in rec["departures"].items())
        day_marks = ["?" if rec["days"][d] is None else ("✓" if rec["days"][d] else "–") for d in DAYS]
        note = "" if rec["daysChecked"] == 7 else f"{7 - rec['daysChecked']} day(s) could not be read; re-run to fill"
        ws.append([rec["originCity"], rec["origin"], rec["destCity"], rec["dest"], rec["destCountry"],
                   None, airlines or "—", *day_marks, None, None, deps or "—", rec["weekStart"], note])
        dl, dr = get_column_letter(first_day), get_column_letter(last_day)
        n = f"{get_column_letter(c_days)}{r}"
        # An exact comparison, not COUNTIF: in COUNTIF a bare "?" is a
        # one-character wildcard and matched every day cell, which marked
        # every route Incomplete. The "~?" escape is not honoured everywhere.
        unknown = f'SUMPRODUCT(--({dl}{r}:{dr}{r}="?"))'
        ws.cell(r, c_days, f'=COUNTIF({dl}{r}:{dr}{r},"✓")')
        ws.cell(r, c_direct, f'=IF({n}>0,"Yes",IF({unknown}>0,"Unknown","No"))')
        ws.cell(r, c_freq, f'=IF({unknown}>0,"Incomplete",IF({n}=7,"Daily",IF({n}>=5,"5-6x weekly",'
                           f'IF({n}>=3,"Alternate days (3-4x weekly)",IF({n}>=1,"1-2x weekly","No direct")))))')

    for cell in ws[1]:
        cell.font, cell.fill = bold, head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = font
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in (c_direct + 1, cols.index("Non-stop departures") + 1))
        for c in range(first_day, last_day + 1):
            row[c - 1].alignment = Alignment(horizontal="center", vertical="top")
    widths = {"Origin": 12, "From": 6, "Destination": 16, "To": 6, "Country": 14, "Direct": 8,
              "Airlines (days)": 46, "Days / week": 8, "Frequency": 24, "Non-stop departures": 40,
              "Week checked": 12, "Note": 30}
    for i, name in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(name, 5)
    last = ws.max_row
    rng = f"{get_column_letter(c_direct)}2:{get_column_letter(c_direct)}{last}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Yes"'], fill=green))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"No"'], fill=grey))
    days_rng = f"{get_column_letter(first_day)}2:{get_column_letter(last_day)}{last}"
    ws.conditional_formatting.add(days_rng, CellIsRule(operator="equal", formula=['"✓"'], fill=green))
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"

    fl = wb.create_sheet("Flights")
    fl.append(["Origin", "From", "Destination", "To", "Date", "Day", "Airline", "Departs", "Arrives",
               "Duration (min)", "Aircraft"])
    for rec in sorted(records, key=lambda r: (r["originCity"], r["destCity"])):
        for day, flights in sorted(rec["flightsByDay"].items()):
            for f in flights or []:
                fl.append([rec["originCity"], rec["origin"], rec["destCity"], rec["dest"], day,
                           DAYS[date.fromisoformat(day).weekday()], f["airline"], f["depart"], f["arrive"],
                           f["durationMinutes"], f["aircraft"]])
    for cell in fl[1]:
        cell.font, cell.fill = bold, head_fill
    for row in fl.iter_rows(min_row=2):
        for cell in row:
            cell.font = font
    for i, w in enumerate([12, 6, 16, 6, 11, 5, 22, 8, 8, 12, 18], 1):
        fl.column_dimensions[get_column_letter(i)].width = w
    fl.freeze_panes = "A2"
    fl.auto_filter.ref = f"A1:K{max(fl.max_row, 2)}"

    about = wb.create_sheet("About")
    lines = [
        ("Route schedules — non-stop service by weekday", True),
        ("", False),
        ("Source: Google Flights, non-stop one-way economy searches, one for each day of the week checked.", False),
        ("Generated by schedule_probe.py in kaushik-work/fk-flight-finder. Nothing here is typed in by hand.", False),
        ("", False),
        ("Day columns: ✓ = at least one non-stop that day · – = none · ? = Google's page could not be read.", False),
        ("Direct, Days / week and Frequency are formulas over the day columns.", False),
        ("Frequency: Daily = 7 days · 5-6x weekly · Alternate days = 3-4 days · 1-2x weekly · No direct.", False),
        ("Flights sheet: one row per non-stop flight found, for checking any route in detail.", False),
        ("", False),
        ("Limits:", True),
        ("One week only. Seasonal schedules (winter schedule from late October) can differ; re-run for another week.", False),
        ("Airline names are as Google shows them; codeshares may appear under the operating airline only.", False),
        ("Flight numbers are not sent by this Google endpoint, so they are not listed.", False),
    ]
    for text, is_bold in lines:
        about.append([text])
        about.cell(about.max_row, 1).font = Font(name="Arial", size=12 if is_bold else 10, bold=is_bold)
    about.column_dimensions["A"].width = 110

    wb.calculation.fullCalcOnLoad = True
    wb.save(path)


def _take_lock() -> Any:
    """On the droplet, refuse to run beside a fare pass."""
    if not os.path.exists(LOCK):
        return None
    import fcntl
    fh = open(LOCK, "a")  # noqa: SIM115 — held for the life of the process
    try:
        fcntl.flock(fh, fcntl.LOCK_EX | fcntl.LOCK_NB)
    except BlockingIOError:
        sys.exit("A fare pass is running (lock held). Run this on another machine, or between passes.")
    return fh


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--origins", nargs="*", help="Origin slugs; default is all.")
    ap.add_argument("--week-start", help="Monday to sample (YYYY-MM-DD); default is 3+ weeks out.")
    ap.add_argument("--limit", type=int, default=0, help="Stop after N destinations per origin.")
    ap.add_argument("--fresh", action="store_true", help="Ignore saved progress and probe everything again.")
    ap.add_argument("--xlsx-only", action="store_true", help="Rebuild the sheet from schedules.json; no requests.")
    args = ap.parse_args()

    os.makedirs(OUT, exist_ok=True)
    saved: dict[str, Any] = {}
    if os.path.exists(JSON_PATH) and not args.fresh:
        with open(JSON_PATH, encoding="utf-8") as fh:
            saved = {f"{r['origin']}-{r['dest']}": r for r in json.load(fh)["routes"]}

    if args.xlsx_only:
        write_xlsx(list(saved.values()), XLSX_PATH)
        print(f"wrote {XLSX_PATH} ({len(saved)} routes)")
        return 0

    lock = _take_lock()  # noqa: F841 — held until exit
    start = date.fromisoformat(args.week_start) if args.week_start else default_week_start()
    if start.weekday() != 0:
        start -= timedelta(days=start.weekday())
    week = [start + timedelta(days=i) for i in range(7)]

    all_origins, dests = scrape.load_routes()
    origins = [o for o in all_origins if not args.origins or o.slug in args.origins]
    routes = [(o, d) for o in origins for d in dests
              if d.code != o.code and (o.code, d.code) not in scrape.SKIP_ROUTES]
    if args.limit:
        routes = [r for o in origins for r in [x for x in routes if x[0] is o][: args.limit]]

    def done(o: Place, d: Place) -> bool:
        rec = saved.get(f"{o.code}-{d.code}")
        return bool(rec) and rec["weekStart"] == start.isoformat() and rec["daysChecked"] == 7

    todo = [(o, d) for o, d in routes if not done(o, d)]
    print(f"week of {start} | {len(routes)} routes, {len(routes) - len(todo)} already done | "
          f"~{len(todo) * 7} requests, ~{len(todo) * 7 * (DELAY * 1.2 + 1) / 3600:.1f}h", flush=True)

    blocked = False
    try:
        for i, (o, d) in enumerate(todo, 1):
            prev = saved.get(f"{o.code}-{d.code}") or {}
            by_day: dict[str, Any] = dict(prev.get("flightsByDay") or {}) if prev.get("weekStart") == start.isoformat() else {}
            for day in week:
                key = day.isoformat()
                if by_day.get(key) is not None:
                    continue  # read last time
                by_day[key] = probe_day(o.code, d.code, key)
                time.sleep(DELAY + random.uniform(0, DELAY * 0.4))
            rec = summarise(o, d, week, by_day)
            saved[f"{o.code}-{d.code}"] = rec
            _save(saved)
            print(f"   [{i}/{len(todo)}] {o.code}->{d.code}: {rec['frequency']}"
                  + (f" | {'; '.join(rec['airlines'])}" if rec["airlines"] else ""), flush=True)
    except scrape.Blocked:
        blocked = True
        print("   ! Google keeps showing a CAPTCHA. Stopped; progress is saved — re-run later to resume.", flush=True)

    write_xlsx(list(saved.values()), XLSX_PATH)
    print(f"wrote {XLSX_PATH} and {JSON_PATH} ({len(saved)} routes)")
    return 3 if blocked else 0


def _save(saved: dict[str, Any]) -> None:
    tmp = JSON_PATH + ".tmp"
    with open(tmp, "w", encoding="utf-8") as fh:
        json.dump({"generatedAt": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()),
                   "routes": sorted(saved.values(), key=lambda r: (r["origin"], r["dest"]))}, fh, indent=1)
    os.replace(tmp, JSON_PATH)


if __name__ == "__main__":
    raise SystemExit(main())
